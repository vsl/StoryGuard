import asyncio
import io
import os
import unittest
import uuid

from docx import Document

from app.parsing import ChunkingConfig, parse_manuscript, token_spans


class ParsingTest(unittest.TestCase):
    def test_detects_structure_across_supported_formats(self) -> None:
        markdown = parse_manuscript(
            "story.md",
            b"Preface\n\n# Chapter One\nAlice waited.\n\n***\n\nBob arrived.\n\n# Epilogue\nDone.",
        )
        self.assertEqual(
            [(chapter.title, len(chapter.scenes)) for chapter in markdown.chapters],
            [("Chapter One", 2), ("Epilogue", 1)],
        )
        self.assertTrue(markdown.chapters[0].text.startswith("Preface"))

        text = parse_manuscript(
            "story.txt", b"Chapter I\nFirst.\n\n---\n\nSecond.\n\nChapter 2: End\nThird."
        )
        self.assertEqual(
            [chapter.title for chapter in text.chapters],
            ["Chapter I", "Chapter 2: End"],
        )
        self.assertEqual(len(text.chapters[0].scenes), 2)

        contents = parse_manuscript(
            "contents.txt",
            b"Contents\n\nChapter I. First\nChapter II. Second\n\n"
            b"Chapter I.\nBody one.\n\nChapter II.\nBody two.",
        )
        self.assertEqual(
            [chapter.title for chapter in contents.chapters],
            ["Chapter I.", "Chapter II."],
        )
        self.assertEqual(sum(len(chapter.scenes) for chapter in contents.chapters), 2)
        self.assertIn(
            "Ignored 2 duplicate chapter markers", contents.metadata["warnings"][0]
        )

        fallback = parse_manuscript("story.txt", b"No headings here.")
        self.assertEqual(len(fallback.chapters), 1)
        self.assertEqual(
            fallback.metadata["warnings"],
            ["No reliable chapter structure; used one chapter"],
        )

        document = Document()
        document.add_heading("Chapter 1", level=1)
        document.add_paragraph("Before the break.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Alice"
        table.cell(0, 1).text = "Paris"
        document.add_paragraph("***")
        document.add_paragraph("After the break.")
        stream = io.BytesIO()
        document.save(stream)
        docx = parse_manuscript("story.docx", stream.getvalue())
        self.assertEqual(docx.metadata["chapter_detection"], "docx_heading_style")
        self.assertIn("Alice\tParis", docx.chapters[0].text)
        self.assertEqual(len(docx.chapters[0].scenes), 2)

    def test_golden_structure_and_false_heading_controls(self) -> None:
        long_body = " ".join(["narrative"] * 40)
        numbered = "\n\n".join(
            part
            for number in range(1, 4)
            for part in (
                str(number),
                f"Title {number}",
                long_body,
                "#",
                f"Second scene {number}. {long_body}",
            )
        )
        parsed_numbered = parse_manuscript("numbered.txt", numbered.encode())
        self.assertEqual(
            [(chapter.title, len(chapter.scenes)) for chapter in parsed_numbered.chapters],
            [("1", 2), ("2", 2), ("3", 2)],
        )

        roman = "\n\n".join(
            (
                "I. FIRST STORY",
                long_body,
                "II. SECOND STORY",
                long_body,
                "III. THIRD STORY",
                long_body,
            )
        )
        self.assertEqual(
            [chapter.title for chapter in parse_manuscript("roman.txt", roman.encode()).chapters],
            ["I. FIRST STORY", "II. SECOND STORY", "III. THIRD STORY"],
        )

        compact_lists = (
            "Shopping list\n\n1\n\nApples\n\n2\n\nPears\n\n3\n\nBread\n\n"
            "I. FIRST OPTION\n\nII. SECOND OPTION\n\nIII. THIRD OPTION"
        )
        self.assertEqual(len(parse_manuscript("list.txt", compact_lists.encode()).chapters), 1)

        markdown_hash = parse_manuscript(
            "hash.md", b"# Chapter One\nFirst paragraph.\n\n#\n\nStill the same scene."
        )
        self.assertEqual(len(markdown_hash.chapters[0].scenes), 1)

        text_hash = parse_manuscript(
            "hash.txt", b"#\nFront matter.\n\nChapter 1\nFirst.\n\n#\n\nSecond."
        )
        self.assertEqual(len(text_hash.chapters[0].scenes), 2)
        self.assertEqual(
            text_hash.metadata["detected_boundaries"],
            {
                "coordinate_system": "normalized_text_line_1_based",
                "chapters": [4],
                "scenes": [7],
            },
        )

        document = Document()
        for number in range(1, 3):
            document.add_heading(f"Chapter {number}", level=1)
            document.add_paragraph(f"First scene {number}.")
            document.add_paragraph("***")
            document.add_paragraph(f"Second scene {number}.")
        stream = io.BytesIO()
        document.save(stream)
        parsed_docx = parse_manuscript("golden.docx", stream.getvalue())
        self.assertEqual(
            [(chapter.title, len(chapter.scenes)) for chapter in parsed_docx.chapters],
            [("Chapter 1", 2), ("Chapter 2", 2)],
        )

    def test_chunk_overlap_offsets_limits_and_docx_security(self) -> None:
        manuscript = " ".join(f"word{i}." for i in range(1_500))
        baseline = parse_manuscript(
            "story.txt", manuscript.encode(), ChunkingConfig(overlap_tokens=0)
        )
        candidate = parse_manuscript(
            "story.txt", manuscript.encode(), ChunkingConfig(overlap_tokens=100)
        )
        baseline_chunks = baseline.chapters[0].scenes[0].chunks
        candidate_chunks = candidate.chapters[0].scenes[0].chunks
        self.assertTrue(all(chunk.token_count <= 900 for chunk in candidate_chunks))
        self.assertGreater(
            sum(chunk.token_count for chunk in candidate_chunks),
            sum(chunk.token_count for chunk in baseline_chunks),
        )
        first, second = candidate_chunks[:2]
        self.assertEqual(
            [first.text[start:end] for start, end in token_spans(first.text)][-100:],
            [second.text[start:end] for start, end in token_spans(second.text)][:100],
        )
        chapter_text = candidate.chapters[0].text
        self.assertTrue(
            all(
                chapter_text[chunk.start_offset : chunk.end_offset] == chunk.text
                for chunk in candidate_chunks
            )
        )

        with self.assertRaisesRegex(ValueError, "Malformed DOCX"):
            parse_manuscript("broken.docx", b"not a zip")
        stream = io.BytesIO()
        document = Document()
        document.add_paragraph("safe text")
        document.save(stream)
        with self.assertRaisesRegex(ValueError, "safe limit"):
            parse_manuscript(
                "large.docx", stream.getvalue(), docx_max_uncompressed_bytes=1
            )


RUN_DATABASE_TESTS = os.environ.get("RUN_DATABASE_TESTS") == "1"

if RUN_DATABASE_TESTS:
    from sqlalchemy import func, select

    from app.db.models.job_run import JobRun
    from app.db.models.manuscript_version import ManuscriptVersion
    from app.db.models.narrative import Chapter, Chunk, Scene
    from app.db.models.project import Project
    from app.db.session import SessionLocal, engine
    from app.manuscripts import create_manuscript_version, minio_client
    from app.queue.tasks.ingestion import run_parse_and_ingest


@unittest.skipUnless(RUN_DATABASE_TESTS, "set RUN_DATABASE_TESTS=1")
class ParsingIntegrationTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.project_id = uuid.uuid4()
        self.object_keys: list[str] = []
        async with SessionLocal() as session:
            session.add(Project(id=self.project_id, title="Parsing test"))
            await session.commit()

    async def asyncTearDown(self) -> None:
        client = minio_client()
        bucket = os.environ["MINIO_BUCKET"]
        for key in self.object_keys:
            await asyncio.to_thread(client.remove_object, bucket, key)
        async with SessionLocal() as session:
            project = await session.get(Project, self.project_id)
            if project:
                await session.delete(project)
                await session.commit()
        await engine.dispose()

    async def _version_and_job(
        self, name: str, data: bytes
    ) -> tuple[ManuscriptVersion, JobRun]:
        async with SessionLocal() as session:
            version = await create_manuscript_version(
                session, self.project_id, name, "text/plain", io.BytesIO(data)
            )
            self.object_keys.append(version.object_key)
            job = JobRun(
                job_type="parse_and_ingest_manuscript",
                project_id=self.project_id,
                manuscript_version_id=version.id,
                idempotency_key=f"parse:{version.id}",
            )
            session.add(job)
            await session.commit()
            await session.refresh(job)
            return version, job

    async def test_ingestion_is_version_scoped_idempotent_and_fails_closed(self) -> None:
        version, job = await self._version_and_job(
            "story.txt", b"Chapter 1\nAlice waited.\n\n***\n\nBob arrived."
        )
        await run_parse_and_ingest(job.id)
        await run_parse_and_ingest(job.id)

        async with SessionLocal() as session:
            persisted_job = await session.get(JobRun, job.id)
            persisted_version = await session.get(ManuscriptVersion, version.id)
            chapter_count = await session.scalar(
                select(func.count()).select_from(Chapter).where(
                    Chapter.manuscript_version_id == version.id
                )
            )
            scene_count = await session.scalar(
                select(func.count()).select_from(Scene).where(
                    Scene.chapter_id.in_(
                        select(Chapter.id).where(
                            Chapter.manuscript_version_id == version.id
                        )
                    )
                )
            )
            chunk_count = await session.scalar(
                select(func.count()).select_from(Chunk).where(
                    Chunk.manuscript_version_id == version.id
                )
            )
            project = await session.get(Project, self.project_id)
        self.assertEqual((persisted_job.status, persisted_job.attempts), ("completed", 1))
        self.assertEqual((chapter_count, scene_count, chunk_count), (1, 2, 2))
        self.assertEqual(persisted_version.status, "ready")
        self.assertIsNotNone(persisted_version.ready_at)
        self.assertEqual(persisted_version.parse_metadata["chunking"]["overlap_tokens"], 100)
        self.assertEqual(project.current_manuscript_version_id, version.id)

        bad_version, bad_job = await self._version_and_job("bad.txt", b"\xff")
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            await run_parse_and_ingest(bad_job.id)
        async with SessionLocal() as session:
            bad_version = await session.get(ManuscriptVersion, bad_version.id)
            bad_job = await session.get(JobRun, bad_job.id)
            project = await session.get(Project, self.project_id)
        self.assertEqual((bad_version.status, bad_job.status), ("failed", "failed"))
        self.assertEqual(project.current_manuscript_version_id, version.id)

        scoped_version, scoped_job = await self._version_and_job(
            "scope.txt", b"Chapter 1\nScoped text."
        )
        other_project_id = uuid.uuid4()
        async with SessionLocal() as session:
            session.add(Project(id=other_project_id, title="Other project"))
            await session.commit()
            scoped_job = await session.get(JobRun, scoped_job.id)
            scoped_job.project_id = other_project_id
            await session.commit()
        with self.assertRaisesRegex(ValueError, "scopes do not match"):
            await run_parse_and_ingest(scoped_job.id)
        async with SessionLocal() as session:
            chunk_count = await session.scalar(
                select(func.count()).select_from(Chunk).where(
                    Chunk.manuscript_version_id == scoped_version.id
                )
            )
            other_project = await session.get(Project, other_project_id)
            await session.delete(other_project)
            await session.commit()
        self.assertEqual(chunk_count, 0)


if __name__ == "__main__":
    unittest.main()
